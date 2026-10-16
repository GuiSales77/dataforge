"""Agente Relator: gera o Excel final com abas de resultado, resumo e log."""
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D9D9")


def _sheet_name(name, used):
    base = "".join(ch for ch in str(name) if ch not in '[]:*?/\\')[:31] or "Resultado"
    n, i = base, 2
    while n.lower() in used:
        suf = f"_{i}"
        n = base[: 31 - len(suf)] + suf
        i += 1
    used.add(n.lower())
    return n


def _style(ws, df):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 0:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, 1):
        sample = df[col].astype(str).head(300)
        width = min(60, max(10, len(str(col)) + 2, int(sample.str.len().quantile(0.9)) + 2 if len(sample) else 10))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 22


def export(path, results: dict, meta: dict, log: list, cfg: dict):
    used = set()
    max_tabs = int(cfg.get("max_abas_saida", 40))
    style = cfg.get("estilizar_excel", True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        # Resumo
        rows = [
            ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")),
            ("Comando", meta.get("comando", "")),
            ("Estratégia", meta.get("estrategia", "")),
            ("Modelo", meta.get("modelo", "")),
            ("Confiança", f"{float(meta.get('confianca', 0)):.0%}"),
            ("Planilhas de entrada", ", ".join(meta.get("entradas", []))),
            ("Linhas de entrada", meta.get("linhas_entrada", 0)),
            ("Tempo de execução (s)", round(meta.get("ms", 0) / 1000, 2)),
            ("", ""),
            ("Resultados", "Linhas"),
        ] + [(k, len(v)) for k, v in results.items()]
        resumo = pd.DataFrame(rows, columns=["Item", "Valor"])
        resumo.to_excel(xw, sheet_name=_sheet_name("Resumo", used), index=False)
        ws = xw.sheets["Resumo"]
        if style:
            _style(ws, resumo)
            ws.column_dimensions["A"].width = 26
            ws.column_dimensions["B"].width = 90
        for k, (name, df) in enumerate(results.items()):
            if k >= max_tabs:
                break
            sn = _sheet_name(name, used)
            df.to_excel(xw, sheet_name=sn, index=False)
            if style:
                _style(xw.sheets[sn], df)
        lg = pd.DataFrame([{"Passo": e["passo"], "Operação": e["descricao"], "Status": e["status"],
                            "Linhas antes": e.get("linhas_antes"), "Linhas depois": e.get("linhas_depois"),
                            "Tempo (ms)": e.get("ms"), "Erro": e.get("erro", "")} for e in log])
        sn = _sheet_name("Log", used)
        lg.to_excel(xw, sheet_name=sn, index=False)
        if style:
            _style(xw.sheets[sn], lg)
            xw.sheets[sn].column_dimensions["B"].width = 80
    return path
